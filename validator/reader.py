"""Read the input workbook and normalise each cell before validation.

The reader is read-only: it opens the workbook with read_only=True and never
saves it, so the original file cannot be changed by a run.
"""

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .rules import Issue, parse_date
from .schema import BY_NAME, COLUMN_NAMES, REQUIRED, STATE_CODES

HEADER_SEARCH_ROWS = 10


class StructureError(Exception):
    """The workbook cannot be validated row by row (e.g. a required column is missing)."""


@dataclass
class Row:
    excel_row: int
    raw: dict
    values: dict
    issues: list = field(default_factory=list)


@dataclass
class Workbook:
    path: Path
    sheet_name: str
    header_row: int
    headers: list
    rows: list
    structure_warnings: list


def normalise_header(text):
    if text is None:
        return ""
    return "_".join(str(text).strip().lower().replace("-", " ").split())


def _find_header_row(sheet):
    """Return (row_number, cells) for the first row naming at least half the columns."""
    wanted = set(COLUMN_NAMES)
    for row_number, cells in enumerate(
            sheet.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS, values_only=True), start=1):
        names = {normalise_header(c) for c in cells}
        if len(names & wanted) >= len(wanted) / 2:
            return row_number, list(cells)
    raise StructureError(
        f"no header row found in the first {HEADER_SEARCH_ROWS} rows; "
        f"expected columns such as {', '.join(COLUMN_NAMES[:4])}")


def _is_whole_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) \
        and float(value).is_integer()


def normalise_cell(column, value):
    """Return (normalised_value, issues) for one cell."""
    kind = BY_NAME[column].kind
    issues = []
    if value is None:
        return None, issues
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return None, issues

    if kind in ("text_id", "pan", "gstin"):
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            issues.append(Issue(column, "NUMERIC_CELL",
                                "stored as a number; store identifiers as text so "
                                "leading zeros and formatting are not lost"))
            value = str(int(value)) if _is_whole_number(value) else str(value)
        return str(value).upper(), issues

    if kind in ("mobile", "pincode"):
        # A phone number or Indian pincode has no meaningful leading zero, so a
        # whole-number cell is safe to accept; a fractional one is not.
        if _is_whole_number(value):
            value = str(int(value))
        value = str(value).replace(" ", "").replace("-", "")
        if kind == "mobile":
            if value.startswith("+91") and len(value) == 13:
                value = value[3:]
            elif value.startswith("91") and len(value) == 12:
                value = value[2:]
            elif value.startswith("0") and len(value) == 11:
                value = value[1:]
        return value, issues

    if kind == "email":
        return str(value).lower(), issues

    if kind == "state":
        match = {k.lower(): k for k in STATE_CODES}.get(str(value).lower())
        return (match or str(value)), issues

    if kind == "date":
        return parse_date(value), issues

    if kind == "money":
        if isinstance(value, str):
            cleaned = value.replace(",", "").replace("₹", "").replace("INR", "").strip()
            try:
                return float(cleaned) if "." in cleaned else int(cleaned), issues
            except ValueError:
                return value, issues
        return value, issues

    if kind == "choice":
        if _is_whole_number(value):
            return int(value), issues
        if isinstance(value, str) and value.isdigit():
            return int(value), issues
        return value, issues

    return str(value), issues


def read_workbook(path, sheet_name=None):
    path = Path(path)
    if not path.exists():
        raise StructureError(f"file not found: {path}")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for non-xlsx input
        raise StructureError(f"could not open {path.name} as an .xlsx workbook: {exc}") from exc

    try:
        sheet = wb[sheet_name] if sheet_name else wb.worksheets[0]
    except KeyError as exc:
        raise StructureError(f"sheet '{sheet_name}' not found; sheets are {wb.sheetnames}") from exc

    header_row, header_cells = _find_header_row(sheet)
    headers = [normalise_header(c) for c in header_cells]

    seen, duplicates = set(), set()
    for h in headers:
        if h and h in seen:
            duplicates.add(h)
        seen.add(h)
    if duplicates:
        raise StructureError(f"duplicate column headers: {', '.join(sorted(duplicates))}")

    missing = [c for c in REQUIRED if c not in headers]
    if missing:
        raise StructureError(f"required columns missing: {', '.join(missing)}")

    structure_warnings = [f"unknown column '{h}' ignored" for h in headers
                          if h and h not in COLUMN_NAMES]

    rows = []
    for excel_row, cells in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if cells is None or all(c is None or (isinstance(c, str) and not c.strip())
                                for c in cells):
            continue  # skip blank rows
        raw = {h: (cells[i] if i < len(cells) else None)
               for i, h in enumerate(headers) if h in COLUMN_NAMES}
        values, issues = {}, []
        for column in COLUMN_NAMES:
            value, cell_issues = normalise_cell(column, raw.get(column))
            values[column] = value
            issues.extend(cell_issues)
        rows.append(Row(excel_row=excel_row, raw=raw, values=values, issues=issues))

    wb.close()
    return Workbook(path=path, sheet_name=sheet.title, header_row=header_row,
                    headers=headers, rows=rows, structure_warnings=structure_warnings)
