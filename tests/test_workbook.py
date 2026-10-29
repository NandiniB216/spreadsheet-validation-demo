"""End-to-end tests on generated workbooks: reading, duplicates, reports, safety."""

import ast
import csv
import hashlib
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from validator.cli import main
from validator.reader import StructureError, read_workbook
from validator.schema import COLUMN_NAMES
from validator.validate import validate_rows

ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 1)


@pytest.fixture()
def sample(tmp_path):
    import importlib.util
    spec = importlib.util.spec_from_file_location("make_sample_data",
                                                  ROOT / "data" / "make_sample_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.build(tmp_path / "sample.xlsx", tmp_path / "expected.csv")


def write_xlsx(path, rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.save(path)
    return path


def test_sample_matches_expected_outcomes(sample):
    xlsx, expected_csv = sample
    result = validate_rows(read_workbook(xlsx), today=TODAY)
    by_row = {r.row.excel_row: r for r in result.results}
    with open(expected_csv, encoding="utf-8") as f:
        expected = list(csv.DictReader(f))
    assert len(expected) == len(result.results)
    for e in expected:
        r = by_row[int(e["excel_row"])]
        if e["expected_status"] == "ACCEPTED":
            assert r.status == "ACCEPTED" and not r.warnings, (e, r.errors, r.warnings)
        elif e["expected_status"] == "ACCEPTED_WITH_WARNING":
            assert r.status == "ACCEPTED" and r.warnings, (e, r.warnings)
        else:
            column, code = e["expected_issue"].split(":")
            assert r.status == "REJECTED"
            assert (column, code) in [(i.column, i.code) for i in r.errors], (e, r.errors)


def test_header_row_found_below_a_title_and_blank_rows_skipped(sample):
    workbook = read_workbook(sample[0])
    assert workbook.header_row == 3
    assert all(any(v is not None for v in r.values.values()) for r in workbook.rows)
    assert workbook.structure_warnings == ["unknown column 'notes' ignored"]


def test_headers_are_matched_regardless_of_case_and_spacing(tmp_path):
    headers = [c.replace("_", " ").title() for c in COLUMN_NAMES]
    path = write_xlsx(tmp_path / "h.xlsx", [headers])
    assert read_workbook(path).rows == []


def test_missing_required_column_stops_the_run(tmp_path):
    headers = [c for c in COLUMN_NAMES if c != "pan"]
    path = write_xlsx(tmp_path / "missing.xlsx", [headers])
    with pytest.raises(StructureError, match="required columns missing: pan"):
        read_workbook(path)


def test_duplicate_header_stops_the_run(tmp_path):
    path = write_xlsx(tmp_path / "dup.xlsx", [[*COLUMN_NAMES, "pan"]])
    with pytest.raises(StructureError, match="duplicate column headers: pan"):
        read_workbook(path)


def test_non_xlsx_file_gives_a_clear_error(tmp_path):
    bad = tmp_path / "not-a-workbook.xlsx"
    bad.write_text("this is plain text")
    with pytest.raises(StructureError, match="could not open"):
        read_workbook(bad)


def test_first_occurrence_of_a_duplicate_is_kept(sample):
    result = validate_rows(read_workbook(sample[0]), today=TODAY)
    sup3 = [r for r in result.results if r.row.values["supplier_id"] == "SUP-0003"]
    assert [r.status for r in sup3] == ["ACCEPTED", "REJECTED"]


def test_cli_writes_reports_and_never_changes_the_input(sample, tmp_path, capsys):
    xlsx = sample[0]
    before = hashlib.sha256(xlsx.read_bytes()).hexdigest()
    out = tmp_path / "out"
    exit_code = main([str(xlsx), "--out", str(out)])
    assert exit_code == 1  # the sample deliberately contains rejected rows
    assert hashlib.sha256(xlsx.read_bytes()).hexdigest() == before

    results = load_workbook(out / "sample-results.xlsx")
    assert results.sheetnames == ["Results", "Issues", "Summary"]
    statuses = [row[len(COLUMN_NAMES) + 1] for row in
                results["Results"].iter_rows(min_row=2, values_only=True)]
    with open(out / "sample-accepted.csv", encoding="utf-8") as f:
        accepted_rows = list(csv.DictReader(f))
    assert len(accepted_rows) == statuses.count("ACCEPTED") == 6
    with open(out / "sample-issues.csv", encoding="utf-8") as f:
        assert len(list(csv.DictReader(f))) >= statuses.count("REJECTED") == 15
    assert "Rejected rows:" in capsys.readouterr().out


def test_cli_exit_code_2_for_structure_errors(tmp_path, capsys):
    path = write_xlsx(tmp_path / "x.xlsx", [["something", "else"]])
    assert main([str(path), "--out", str(tmp_path / "o")]) == 2
    assert "STRUCTURE ERROR" in capsys.readouterr().err


def test_cli_exit_code_0_when_every_row_passes(tmp_path):
    good = ["SUP-0100", "Valid Co", "a@valid.example", "9999900100", "XQZCA1001K", None,
            "Delhi", "110001", date(2026, 1, 5), 1000, 30]
    path = write_xlsx(tmp_path / "ok.xlsx", [list(COLUMN_NAMES), good])
    assert main([str(path), "--out", str(tmp_path / "o")]) == 0


def test_package_makes_no_network_calls():
    """Static guard: nothing in the package imports a networking module."""
    banned = {"socket", "requests", "urllib", "http", "httpx", "aiohttp", "smtplib", "ftplib"}
    for source in (ROOT / "validator").glob("*.py"):
        tree = ast.parse(source.read_text())
        for node in ast.walk(tree):
            names = []
            if isinstance(node, ast.Import):
                names = [a.name for a in node.names]
            elif isinstance(node, ast.ImportFrom) and node.module:
                names = [node.module]
            for name in names:
                assert name.split(".")[0] not in banned, f"{source.name} imports {name}"
